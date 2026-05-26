"""
Tests for document ingestion — US-102 (DOCX), US-103 (TXT), Bonus (PDF).
"""

import os
import tempfile
from pathlib import Path

import pytest

from ingestion import ingest, get_reader, IngestionError
from ingestion.base_reader import Document
from ingestion.docx_reader import DocxReader
from ingestion.txt_reader import TxtReader


# ── Helpers ─────────────────────────────────────────────────────────────────

def _create_temp_txt(content: str, encoding: str = "utf-8") -> str:
    """Create a temporary .txt file with the given content."""
    fd, path = tempfile.mkstemp(suffix=".txt")
    with os.fdopen(fd, "w", encoding=encoding) as f:
        f.write(content)
    return path


def _create_temp_docx(paragraphs: list[str], tables: list[list[list[str]]] = None) -> str:
    """Create a temporary .docx file with the given content."""
    from docx import Document as DocxDocument

    doc = DocxDocument()
    for para in paragraphs:
        doc.add_paragraph(para)

    if tables:
        for table_data in tables:
            rows = len(table_data)
            cols = len(table_data[0]) if table_data else 0
            table = doc.add_table(rows=rows, cols=cols)
            for i, row_data in enumerate(table_data):
                for j, cell_text in enumerate(row_data):
                    table.rows[i].cells[j].text = cell_text

    fd, path = tempfile.mkstemp(suffix=".docx")
    os.close(fd)
    doc.save(path)
    return path


# ── US-103: TXT Reader Tests ───────────────────────────────────────────────

class TestTxtReader:
    """Tests for plain text file ingestion."""

    def test_read_basic_txt(self):
        """Test basic text file reading."""
        content = "Hello, World!\nThis is a test document."
        path = _create_temp_txt(content)

        try:
            reader = TxtReader()
            doc = reader.read(path)

            assert isinstance(doc, Document)
            assert "Hello, World!" in doc.text
            assert "test document" in doc.text
            assert doc.metadata["file_type"] == "txt"
            assert doc.metadata["line_count"] == 2
        finally:
            os.unlink(path)

    def test_read_empty_txt(self):
        """Test empty file handling."""
        path = _create_temp_txt("")

        try:
            reader = TxtReader()
            doc = reader.read(path)

            assert doc.text == ""
            assert doc.metadata["char_count"] == 0
            assert doc.metadata["line_count"] == 0
        finally:
            os.unlink(path)

    def test_read_multiline_txt(self):
        """Test multi-line text with varied whitespace."""
        content = "Line 1\n\nLine 3\n   \nLine 5"
        path = _create_temp_txt(content)

        try:
            reader = TxtReader()
            doc = reader.read(path)

            assert doc.metadata["line_count"] == 5
            assert doc.char_count > 0
        finally:
            os.unlink(path)

    def test_read_utf8_with_bom(self):
        """Test BOM stripping from UTF-8 files."""
        content = "\ufeffBOM content here"
        path = _create_temp_txt(content)

        try:
            reader = TxtReader()
            doc = reader.read(path)

            assert not doc.text.startswith("\ufeff")
            assert doc.text.startswith("BOM content here")
        finally:
            os.unlink(path)

    def test_file_not_found(self):
        """Test FileNotFoundError for missing files."""
        reader = TxtReader()
        with pytest.raises(FileNotFoundError):
            reader.read("/nonexistent/path/file.txt")

    def test_unsupported_extension(self):
        """Test IngestionError for wrong extension."""
        fd, path = tempfile.mkstemp(suffix=".csv")
        os.close(fd)

        try:
            reader = TxtReader()
            with pytest.raises(IngestionError):
                reader.read(path)
        finally:
            os.unlink(path)

    def test_metadata_populated(self):
        """Verify all expected metadata fields are present."""
        path = _create_temp_txt("Test content for metadata.")

        try:
            reader = TxtReader()
            doc = reader.read(path)

            required_keys = [
                "source", "filename", "file_extension",
                "file_size_bytes", "file_type", "encoding",
                "line_count", "char_count",
            ]
            for key in required_keys:
                assert key in doc.metadata, f"Missing metadata key: {key}"
        finally:
            os.unlink(path)


# ── US-102: DOCX Reader Tests ──────────────────────────────────────────────

class TestDocxReader:
    """Tests for Word document ingestion."""

    def test_read_paragraphs(self):
        """Test extraction of paragraph text."""
        paragraphs = ["First paragraph.", "Second paragraph.", "Third paragraph."]
        path = _create_temp_docx(paragraphs)

        try:
            reader = DocxReader()
            doc = reader.read(path)

            assert isinstance(doc, Document)
            for para in paragraphs:
                assert para in doc.text
            assert doc.metadata["paragraph_count"] == 3
            assert doc.metadata["file_type"] == "docx"
        finally:
            os.unlink(path)

    def test_read_with_tables(self):
        """Test extraction of table cell text."""
        paragraphs = ["Document with a table."]
        tables = [[
            ["Header 1", "Header 2"],
            ["Cell A", "Cell B"],
        ]]
        path = _create_temp_docx(paragraphs, tables)

        try:
            reader = DocxReader()
            doc = reader.read(path)

            assert "Header 1" in doc.text
            assert "Cell A" in doc.text
            assert doc.metadata["table_count"] == 1
        finally:
            os.unlink(path)

    def test_read_empty_docx(self):
        """Test handling of empty DOCX file."""
        path = _create_temp_docx([])

        try:
            reader = DocxReader()
            doc = reader.read(path)

            assert doc.text.strip() == ""
            assert doc.metadata["paragraph_count"] == 0
        finally:
            os.unlink(path)

    def test_metadata_populated(self):
        """Verify all expected DOCX metadata fields are present."""
        path = _create_temp_docx(["Test paragraph."])

        try:
            reader = DocxReader()
            doc = reader.read(path)

            required_keys = [
                "source", "filename", "file_type",
                "paragraph_count", "table_count",
            ]
            for key in required_keys:
                assert key in doc.metadata, f"Missing metadata key: {key}"
        finally:
            os.unlink(path)


# ── Auto-detect Tests ───────────────────────────────────────────────────────

class TestAutoDetect:
    """Test the auto-detect reader factory."""

    def test_detect_txt(self):
        """Verify TxtReader is returned for .txt files."""
        reader = get_reader("document.txt")
        assert isinstance(reader, TxtReader)

    def test_detect_docx(self):
        """Verify DocxReader is returned for .docx files."""
        reader = get_reader("document.docx")
        assert isinstance(reader, DocxReader)

    def test_unsupported_format(self):
        """Verify IngestionError for unsupported formats."""
        with pytest.raises(IngestionError):
            get_reader("document.xlsx")

    def test_ingest_convenience(self):
        """Test the convenience ingest() function end-to-end."""
        content = "Convenience function test."
        path = _create_temp_txt(content)

        try:
            doc = ingest(path)
            assert "Convenience function test" in doc.text
        finally:
            os.unlink(path)
